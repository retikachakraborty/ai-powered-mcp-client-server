"""Safe document storage, extraction, search, and editing helpers."""

from __future__ import annotations

import os
import re
import sys
import tempfile
from collections.abc import Callable, Iterator
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class DocumentError(ValueError):
    """A user-facing document operation error."""


SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx"})
EDITABLE_EXTENSIONS = frozenset({".txt", ".md"})
DEFAULT_MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
DEFAULT_MAX_EDIT_REPLACEMENTS = 100


def _positive_environment_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if value is None:
        return default
    try:
        parsed = int(value)
    except ValueError:
        return default
    return parsed if parsed > 0 else default


MAX_DOCUMENT_BYTES = _positive_environment_int(
    "MCP_MAX_DOCUMENT_BYTES", DEFAULT_MAX_DOCUMENT_BYTES
)
MAX_EDIT_REPLACEMENTS = _positive_environment_int(
    "MCP_MAX_EDIT_REPLACEMENTS", DEFAULT_MAX_EDIT_REPLACEMENTS
)


def _document_id_is_filename(doc_id: str) -> bool:
    if not isinstance(doc_id, str) or not doc_id or "\x00" in doc_id:
        return False
    if doc_id in {".", ".."} or Path(doc_id).name != doc_id:
        return False
    # Reject Windows separators even when the server runs on POSIX.
    return "/" not in doc_id and "\\" not in doc_id


def resolve_document_path(
    documents_dir: Path, doc_id: str, *, require_exists: bool = True
) -> Path:
    """Resolve a document filename without allowing traversal or symlinks."""
    documents_dir = documents_dir.resolve()
    if not _document_id_is_filename(doc_id):
        raise DocumentError("Document ID must be a filename, not a path.")

    candidate = documents_dir / doc_id
    if candidate.is_symlink():
        raise DocumentError("Symbolic-link documents are not supported.")
    path = candidate.resolve(strict=False)
    try:
        path.relative_to(documents_dir)
    except ValueError as exc:
        raise DocumentError("Document must be inside the documents directory.") from exc

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise DocumentError("Unsupported document format.")
    if require_exists and not path.is_file():
        raise DocumentError(f"Document '{doc_id}' not found.")
    return path


def iter_document_paths(documents_dir: Path) -> Iterator[Path]:
    """Yield supported regular documents in stable filename order."""
    if not documents_dir.is_dir():
        return
    for path in sorted(documents_dir.iterdir(), key=lambda item: item.name.casefold()):
        if (
            path.is_file()
            and not path.is_symlink()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ):
            yield path


def _read_text_file(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentError(f"Document '{path.name}' is not valid UTF-8 text.") from exc
    except OSError as exc:
        raise DocumentError(f"Could not read document '{path.name}'.") from exc


def _ensure_file_size(path: Path, max_bytes: int | None = None) -> None:
    limit = MAX_DOCUMENT_BYTES if max_bytes is None else max_bytes
    if limit < 1:
        raise DocumentError("The document size limit must be positive.")
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise DocumentError(f"Could not inspect document '{path.name}'.") from exc
    if size > limit:
        raise DocumentError(
            f"Document '{path.name}' exceeds the configured size limit of {limit} bytes."
        )


def _extract_pdf(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise DocumentError(f"Could not extract text from PDF '{path.name}'.") from exc
    content = "\n".join(pages).strip()
    if not content and reader.pages:
        return "[No extractable text found; this PDF may contain scanned images.]"
    return content


def _iter_docx_text(document: Document) -> Iterator[str]:
    for paragraph in document.paragraphs:
        if paragraph.text:
            yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                yield " | ".join(cells)


def _extract_docx(path: Path) -> str:
    try:
        document = Document(str(path))
        return "\n".join(_iter_docx_text(document))
    except Exception as exc:
        raise DocumentError(f"Could not extract text from DOCX '{path.name}'.") from exc


def read_document_content(
    documents_dir: Path, doc_id: str, *, max_bytes: int | None = None
) -> str:
    path = resolve_document_path(documents_dir, doc_id)
    _ensure_file_size(path, max_bytes)
    extension = path.suffix.lower()
    if extension in {".txt", ".md"}:
        return _read_text_file(path)
    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path)
    raise DocumentError("Unsupported document format.")


def list_documents(documents_dir: Path) -> list[str]:
    return [path.name for path in iter_document_paths(documents_dir)]


def _atomic_write_text(path: Path, content: str) -> None:
    mode = path.stat().st_mode & 0o777
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            temporary.write(content)
            temporary.flush()
            os.fsync(temporary.fileno())
        os.chmod(temporary_path, mode)
        os.replace(temporary_path, path)
    except OSError as exc:
        raise DocumentError(f"Could not safely write document '{path.name}'.") from exc
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass


def replace_document_text(
    documents_dir: Path,
    doc_id: str,
    old_str: str,
    new_str: str,
    *,
    expected_replacements: int | None = None,
    max_replacements: int | None = None,
) -> int:
    path = resolve_document_path(documents_dir, doc_id)
    if path.suffix.lower() not in EDITABLE_EXTENSIONS:
        raise DocumentError("Only TXT and Markdown documents can be edited.")
    if not old_str:
        raise DocumentError("The text to replace cannot be empty.")
    _ensure_file_size(path)
    content = _read_text_file(path)
    replacements = content.count(old_str)
    if replacements == 0:
        raise DocumentError("The specified text was not found in the document.")
    if expected_replacements is not None and replacements != expected_replacements:
        raise DocumentError(
            f"Expected {expected_replacements} replacement(s), found {replacements}; "
            "the document was not changed."
        )
    replacement_limit = (
        MAX_EDIT_REPLACEMENTS if max_replacements is None else max_replacements
    )
    if replacement_limit < 1:
        raise DocumentError("The replacement limit must be positive.")
    if replacements > replacement_limit:
        raise DocumentError(
            f"Replacement count {replacements} exceeds the configured limit of "
            f"{replacement_limit}; the document was not changed."
        )
    _atomic_write_text(path, content.replace(old_str, new_str))
    return replacements


def _snippet(content: str, query: str, radius: int = 120) -> str:
    match = re.search(re.escape(query), content, flags=re.IGNORECASE)
    if match is None:
        return content[: radius * 2].strip()
    start = max(0, match.start() - radius)
    end = min(len(content), match.end() + radius)
    prefix = "…" if start else ""
    suffix = "…" if end < len(content) else ""
    return f"{prefix}{content[start:end].strip()}{suffix}"


def search_documents(
    documents_dir: Path,
    query: str,
    *,
    limit: int = 20,
    on_skip: Callable[[str, str], None] | None = None,
) -> list[dict[str, str | int]]:
    """Search extracted document text and return ranked snippets."""
    query = query.strip()
    if not query:
        raise DocumentError("Search query cannot be empty.")
    if limit < 1 or limit > 100:
        raise DocumentError("Search limit must be between 1 and 100.")

    folded_query = query.casefold()
    matches: list[dict[str, str | int]] = []
    for path in iter_document_paths(documents_dir):
        try:
            content = read_document_content(documents_dir, path.name)
        except Exception as error:
            message = f"Skipped document '{path.name}' ({type(error).__name__})."
            if on_skip is not None:
                on_skip(path.name, type(error).__name__)
            else:
                print(f"[search] {message}", file=sys.stderr)
            continue
        score = content.casefold().count(folded_query)
        if score:
            matches.append(
                {
                    "document": path.name,
                    "matches": score,
                    "snippet": _snippet(content, query),
                }
            )
    matches.sort(key=lambda item: (-int(item["matches"]), str(item["document"]).casefold()))
    return matches[:limit]
