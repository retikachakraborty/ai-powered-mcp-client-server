from pathlib import Path

import pytest
from docx import Document
from reportlab.pdfgen import canvas

import core.documents as document_module
from core.documents import (
    DocumentError,
    list_documents,
    read_document_content,
    replace_document_text,
    resolve_document_path,
    search_documents,
)


def _make_pdf(path: Path, text: str) -> None:
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, text)
    pdf.save()


def _make_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("DOCX paragraph")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "First"
    table.rows[0].cells[1].text = "Second"
    document.save(path)


def test_extracts_all_supported_formats(tmp_path: Path) -> None:
    (tmp_path / "plain.txt").write_text("Plain text", encoding="utf-8")
    (tmp_path / "notes.md").write_text("# Markdown", encoding="utf-8")
    _make_pdf(tmp_path / "report.pdf", "PDF content")
    _make_docx(tmp_path / "word.docx")

    assert list_documents(tmp_path) == ["notes.md", "plain.txt", "report.pdf", "word.docx"]
    assert read_document_content(tmp_path, "plain.txt") == "Plain text"
    assert read_document_content(tmp_path, "notes.md") == "# Markdown"
    assert "PDF content" in read_document_content(tmp_path, "report.pdf")
    docx_content = read_document_content(tmp_path, "word.docx")
    assert "DOCX paragraph" in docx_content
    assert "First | Second" in docx_content


def test_rejects_traversal_unsupported_files_and_symlinks(tmp_path: Path) -> None:
    (tmp_path / "safe.txt").write_text("safe", encoding="utf-8")
    (tmp_path / "secret.txt").write_text("secret", encoding="utf-8")
    (tmp_path / "link.txt").symlink_to(tmp_path / "secret.txt")

    for doc_id in ("../secret.txt", "nested/safe.txt", "safe.py", "link.txt"):
        with pytest.raises(DocumentError):
            resolve_document_path(tmp_path, doc_id)


def test_edit_is_exact_and_persistent(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("alpha\nalpha\n", encoding="utf-8")

    assert replace_document_text(tmp_path, "notes.md", "alpha", "beta") == 2
    assert path.read_text(encoding="utf-8") == "beta\nbeta\n"
    with pytest.raises(DocumentError, match="not found"):
        replace_document_text(tmp_path, "notes.md", "missing", "value")


def test_edit_expected_count_and_limit_prevent_unintended_replacements(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("alpha alpha", encoding="utf-8")

    with pytest.raises(DocumentError, match="Expected 1 replacement"):
        replace_document_text(
            tmp_path, "notes.md", "alpha", "beta", expected_replacements=1
        )
    assert path.read_text(encoding="utf-8") == "alpha alpha"

    with pytest.raises(DocumentError, match="exceeds the configured limit"):
        replace_document_text(
            tmp_path, "notes.md", "alpha", "beta", max_replacements=1
        )
    assert path.read_text(encoding="utf-8") == "alpha alpha"


def test_document_size_limit_is_inclusive_at_boundary(tmp_path: Path) -> None:
    path = tmp_path / "small.txt"
    path.write_text("12345", encoding="utf-8")

    assert read_document_content(tmp_path, "small.txt", max_bytes=5) == "12345"
    with pytest.raises(DocumentError, match="size limit"):
        read_document_content(tmp_path, "small.txt", max_bytes=4)


def test_search_reports_skipped_documents_without_changing_result_schema(
    tmp_path: Path, monkeypatch
) -> None:
    (tmp_path / "good.txt").write_text("hey", encoding="utf-8")
    (tmp_path / "large.txt").write_text("hey too large", encoding="utf-8")
    monkeypatch.setattr(document_module, "MAX_DOCUMENT_BYTES", 5)
    skipped = []

    results = search_documents(
        tmp_path,
        "hey",
        on_skip=lambda name, reason: skipped.append((name, reason)),
    )

    assert results == [{"document": "good.txt", "matches": 1, "snippet": "hey"}]
    assert skipped == [("large.txt", "DocumentError")]


def test_search_is_case_insensitive_ranked_and_snippet_bounded(tmp_path: Path) -> None:
    (tmp_path / "many.txt").write_text("needle needle and more text", encoding="utf-8")
    (tmp_path / "one.md").write_text("A Needle appears here", encoding="utf-8")
    (tmp_path / "none.txt").write_text("unrelated", encoding="utf-8")

    results = search_documents(tmp_path, "NEEDLE")

    assert [result["document"] for result in results] == ["many.txt", "one.md"]
    assert results[0]["matches"] == 2
    assert "needle" in str(results[0]["snippet"]).casefold()
    with pytest.raises(DocumentError, match="cannot be empty"):
        search_documents(tmp_path, " ")


def test_search_skips_unreadable_or_invalid_documents(tmp_path: Path, capsys) -> None:
    (tmp_path / "good.txt").write_text("find me", encoding="utf-8")
    (tmp_path / "bad.txt").write_bytes(b"\xff\xfe")

    results = search_documents(tmp_path, "find")

    assert [result["document"] for result in results] == ["good.txt"]
    diagnostic = capsys.readouterr().err
    assert "Skipped document 'bad.txt' (DocumentError)." in diagnostic
    assert "\xff" not in diagnostic
